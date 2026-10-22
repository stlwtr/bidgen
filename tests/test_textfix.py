from bidgen import textfix


def test_kangxi_radicals_repaired():
    assert textfix.fix("投标⼈") == "投标人"
    assert textfix.fix("⽂件") == "文件"
    assert textfix.fix("第⼀章") == "第一章"
    assert textfix.fix("⽬的") == "目的"


def test_radicals_supplement_repaired():
    assert textfix.fix("⿊⻰江") == "黑龙江"
    assert textfix.fix("第⼀⻚") == "第一页"
    assert textfix.fix("⻔⻋⻅") == "门车见"
    assert textfix.fix("⻓⻛⻜") == "长风飞"
    assert textfix.fix("⻢⻨⻩") == "马麦黄"


def test_door_radical_simplified():
    # U+2F3E NFKC -> 戶, but the simplified glyph is intended.
    assert textfix.fix("⼾") == "户"
    assert textfix.fix("账⼾") == "账户"


def test_plain_text_untouched():
    assert textfix.fix("正常文字ABC123") == "正常文字ABC123"
    assert textfix.fix("") == ""


def test_fix_docx_repairs_only_glyphs(tmp_path):
    from docx import Document

    doc = Document()
    doc.add_paragraph("投标⼈⽂件⿊⻰江")  # radicals in a real document
    doc.add_paragraph("全角标点：（1）保持不变")  # must NOT be folded
    path = tmp_path / "glyphs.docx"
    doc.save(str(path))

    textfix.fix_docx(str(path))

    reopened = Document(str(path))
    assert reopened.paragraphs[0].text == "投标人文件黑龙江"
    assert reopened.paragraphs[1].text == "全角标点：（1）保持不变"
