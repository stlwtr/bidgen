"""docx_builder 单元测试：当前锁定「其他资料」章节只保留标题的决策。"""

from __future__ import annotations

from types import SimpleNamespace

from docx import Document

from bidgen.docx_builder import _materials_section
from bidgen.style import StyleSpec


def test_materials_section_keeps_only_titles():
    """格式十六「其他资料」：只输出两个标题，不生成分类小标题与表格。"""
    doc = Document()
    _materials_section(doc, SimpleNamespace(), StyleSpec())

    text = "\n".join(p.text for p in doc.paragraphs)
    assert "（一）招标文件要求提供的其他资料" in text
    assert "（二）投标人认为需提供的其他资料" in text
    # 分类小标题与附表不再生成（第四、五章已编入审查/评分表格，此处不重复）
    assert "资格审查" not in text
    assert len(doc.tables) == 0
