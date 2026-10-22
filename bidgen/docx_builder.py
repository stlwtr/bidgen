"""Build the bid-document DOCX from a parsed tender template and extracted data."""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from typing import Dict, List, Optional, Sequence

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from . import config, extract, fill, punctuation, structure
from .extract import Section
from .style import StyleSpec

_TOC_PREFIX_RE = re.compile(r"^[一二三四五六七八九十]+[、.]\s*")

# 位置 = 投标文件目录序号
#   1=投标承诺书 2=资格承诺函 3=授权委托书 5=主要商务要求承诺书
# 所有章节仅遵照招标文件模板原文输出，不附加自造的头部、签署块或提示段。
_DEVIATION_POSITION = 6
_MATERIALS_POSITION = 16

# Alignment detected by structure (from the PDF position) → Word alignment.
_ALIGN_BY_NAME = {
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
}


@dataclass
class BidContext:
    template: structure.BidTemplate
    fields: Dict[str, Optional[str]]
    supplier: Dict[str, str]
    qualification: Section
    compliance: Section
    scoring: Section
    purchase: Section
    tech: Section
    spec: StyleSpec = field(default_factory=StyleSpec)


def _txt(text: str) -> str:
    return punctuation.normalize(text) if text else text


# --------------------------------------------------------------------------- #
# Low-level styling helpers
# --------------------------------------------------------------------------- #
def _style_run(run, spec: StyleSpec, font: str, size: float, bold: bool = False) -> None:
    run.font.name = spec.latin_font
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), spec.latin_font)
    rfonts.set(qn("w:hAnsi"), spec.latin_font)
    rfonts.set(qn("w:eastAsia"), font)


def _set_first_line_indent(paragraph, chars: int, size: float) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    ind = ppr.get_or_add_ind()
    ind.set(qn("w:firstLineChars"), str(chars * 100))
    ind.set(qn("w:firstLine"), str(int(size * 20 * chars)))  # twips fallback


def _add_runs(paragraph, text: str, spec: StyleSpec, font: str, size: float, bold: bool) -> None:
    """Write *text* as runs, underlining the spans marked by fill.split_markup."""
    for segment, underlined in fill.split_markup(text):
        if segment == "":
            continue
        run = paragraph.add_run(_txt(segment))
        _style_run(run, spec, font, size, bold)
        if underlined:
            run.font.underline = True


def _add_paragraph(
    doc,
    text: str,
    spec: StyleSpec,
    *,
    style: Optional[str] = None,
    font: Optional[str] = None,
    size: Optional[float] = None,
    bold: bool = False,
    align=WD_ALIGN_PARAGRAPH.JUSTIFY,
    indent_chars: int = 0,
    line_spacing: Optional[float] = None,
    space_before: float = 0,
    space_after: float = 0,
    page_break_before: bool = False,
):
    paragraph = doc.add_paragraph(style=style)
    _add_runs(paragraph, text, spec, font or spec.body_font, size or spec.body_size, bold)
    pf = paragraph.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    if line_spacing is not None:
        pf.line_spacing = line_spacing
    pf.alignment = align if align is not None else WD_ALIGN_PARAGRAPH.JUSTIFY
    if page_break_before:
        pf.page_break_before = True
    if indent_chars:
        _set_first_line_indent(paragraph, indent_chars, size or spec.body_size)
    return paragraph


def _add_heading(doc, text: str, level: int, spec: StyleSpec, *, numbered: bool,
                 page_break_before: bool = False):
    size = {1: spec.h1_size, 2: spec.h2_size, 3: spec.h3_size}.get(level, spec.body_size)
    style_name = f"Heading {level}"
    paragraph = doc.add_paragraph(style=style_name)
    run = paragraph.add_run(_txt(text))
    _style_run(run, spec, spec.heading_font, size, bold=True)
    pf = paragraph.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf.line_spacing = spec.heading_line_spacing
    pf.space_before = Pt(12 if level == 1 else 6)
    pf.space_after = Pt(6)
    if level == 1 or page_break_before:
        pf.page_break_before = True
    return paragraph


def _placed_tables(fmt: Optional[structure.Format], position: int):
    """Template tables paired with their (page, top) position, in document order."""
    if fmt is None or not fmt.tables:
        return []
    if len(fmt.table_spans) == len(fmt.tables):
        pairs = list(zip(fmt.table_spans, fmt.tables))
    else:
        pairs = [((0, 0.0), rows) for rows in fmt.tables]
    pairs.sort(key=lambda pair: pair[0])
    return pairs


def _add_table(doc, rows: Sequence[Sequence[str]], spec: StyleSpec) -> None:
    if not rows:
        return
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=0, cols=cols)
    table.style = "Table Grid"
    for row in rows:
        cells = table.add_row().cells
        for i in range(cols):
            value = (row[i] if i < len(row) else "").replace("\n", "")
            cell = cells[i]
            cell.text = ""
            paragraph = cell.paragraphs[0]
            run = paragraph.add_run(_txt(value))
            _style_run(run, spec, spec.body_font, spec.table_size)
            paragraph.paragraph_format.line_spacing = 1.0
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)


def _add_field(paragraph, instr_text: str, spec: StyleSpec) -> None:
    run = paragraph.add_run()
    _style_run(run, spec, spec.body_font, spec.header_size)
    r = run._element
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instr_text
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    r.append(begin)
    r.append(instr)
    r.append(end)


# --------------------------------------------------------------------------- #
# Page / style setup
# --------------------------------------------------------------------------- #
def _init_styles(doc, spec: StyleSpec) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = spec.latin_font
    normal.font.size = Pt(spec.body_size)
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), spec.body_font)
    rfonts.set(qn("w:ascii"), spec.latin_font)
    rfonts.set(qn("w:hAnsi"), spec.latin_font)

    for level, size in ((1, spec.h1_size), (2, spec.h2_size), (3, spec.h3_size)):
        style = doc.styles[f"Heading {level}"]
        style.font.name = spec.latin_font
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        srpr = style.element.get_or_add_rPr()
        srf = srpr.get_or_add_rFonts()
        srf.set(qn("w:eastAsia"), spec.heading_font)
        srf.set(qn("w:ascii"), spec.latin_font)
        srf.set(qn("w:hAnsi"), spec.latin_font)


def _apply_layout(section, spec: StyleSpec) -> None:
    section.top_margin = Cm(spec.margin_top_cm)
    section.bottom_margin = Cm(spec.margin_bottom_cm)
    section.left_margin = Cm(spec.margin_left_cm)
    section.right_margin = Cm(spec.margin_right_cm)
    section.header_distance = Cm(spec.header_distance_cm)
    section.footer_distance = Cm(spec.footer_distance_cm)


def _setup_heading_numbering(doc, spec: StyleSpec) -> bool:
    """Bind a Chinese '一、二、' multilevel list to Heading 1. Returns success."""
    try:
        numbering = doc.part.numbering_part.element
    except (NotImplementedError, KeyError, AttributeError):
        return False
    try:
        existing = [int(e.get(qn("w:abstractNumId"))) for e in numbering.findall(qn("w:abstractNum"))]
        abs_id = (max(existing) + 1) if existing else 1
        nums = [int(e.get(qn("w:numId"))) for e in numbering.findall(qn("w:num"))]
        num_id = (max(nums) + 1) if nums else 1

        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abs_id))
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), "0")
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        fmt = OxmlElement("w:numFmt")
        fmt.set(qn("w:val"), "chineseCounting")
        text = OxmlElement("w:lvlText")
        text.set(qn("w:val"), "%1、")
        jc = OxmlElement("w:lvlJc")
        jc.set(qn("w:val"), "left")
        for el in (start, fmt, text, jc):
            lvl.append(el)
        abstract.append(lvl)
        numbering.insert(0, abstract)

        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(num_id))
        ref = OxmlElement("w:abstractNumId")
        ref.set(qn("w:val"), str(abs_id))
        num.append(ref)
        numbering.append(num)

        style = doc.styles["Heading 1"]
        ppr = style.element.get_or_add_pPr()
        numpr = OxmlElement("w:numPr")
        ilvl = OxmlElement("w:ilvl")
        ilvl.set(qn("w:val"), "0")
        nid = OxmlElement("w:numId")
        nid.set(qn("w:val"), str(num_id))
        numpr.append(ilvl)
        numpr.append(nid)
        ppr.append(numpr)
        return True
    except Exception:
        return False


def _set_header_footer(section, spec: StyleSpec) -> None:
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False

    header = section.header.paragraphs[0]
    header.text = ""
    run = header.add_run(spec.header_text)
    _style_run(run, spec, spec.body_font, spec.header_size)
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    ppr = header._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "auto")
    pbdr.append(bottom)
    ppr.append(pbdr)

    footer = section.footer.paragraphs[0]
    footer.text = ""
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = footer.add_run("第 ")
    _style_run(r1, spec, spec.body_font, spec.header_size)
    _add_field(footer, " PAGE ", spec)
    r2 = footer.add_run(" 页 共 ")
    _style_run(r2, spec, spec.body_font, spec.header_size)
    _add_field(footer, " NUMPAGES ", spec)
    r3 = footer.add_run(" 页")
    _style_run(r3, spec, spec.body_font, spec.header_size)


def _add_toc(doc, spec: StyleSpec) -> None:
    # A plain (non-heading) title so the TOC has no entry for itself and the
    # automatic chapter numbering does not touch it.
    _add_paragraph(doc, "投标文件目录", spec, font=spec.heading_font, size=spec.h1_size,
                   bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=12, space_after=6)
    paragraph = doc.add_paragraph()
    run = paragraph.add_run()
    _style_run(run, spec, spec.body_font, spec.body_size)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = r'TOC \o "1-3" \h \z \u'
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = _txt("（打开文档后按 F9 更新域生成目录）")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for el in (begin, instr, separate, placeholder, end):
        run._element.append(el)


# --------------------------------------------------------------------------- #
# Content blocks
# --------------------------------------------------------------------------- #
def _deviation_table(doc, ctx: BidContext, spec: StyleSpec,
                     template: Optional[structure.Format] = None) -> None:
    """技术偏离表从招标文件提供的模板提取。

    When the tender ships its own deviation-table template, it is already
    rendered in-place by `_placed_tables` (between the centered title and the
    说明 block); nothing more is needed here.  Only when the tender provides no
    such table do we fall back to a blank header skeleton.
    """
    if template is not None and template.tables:
        return
    header = ["序号", "标的名称", "招标技术要求", "投标人提供的响应内容", "偏离程度", "佐证文件名称及所在页码"]
    rows: List[List[str]] = [header, ["", "", "", "", "", ""], ["", "", "", "", "", ""],
                             ["", "", "", "", "", ""]]
    _add_table(doc, rows, spec)


def _materials_section(doc, ctx: BidContext, spec: StyleSpec) -> None:
    """各类证明材料（格式十六）：(一) 按类别附 资格审查/符合性审查/评分项 表格。"""
    tech, biz = extract.split_scoring(ctx.scoring)
    groups = [
        ("1.资格审查", ctx.qualification),
        ("2.符合性审查", ctx.compliance),
        ("3.技术评审", tech),
        ("4.商务评审", biz),
    ]
    _add_paragraph(doc, "（一）招标文件要求提供的其他资料", spec, bold=True, space_after=4)
    for title, section in groups:
        _add_paragraph(doc, title, spec, space_before=2, space_after=2)
        if section is not None and not section.is_empty():
            for table_rows in section.tables:
                _add_table(doc, table_rows, spec)
                _add_paragraph(doc, "", spec, space_after=4)
    _add_paragraph(doc, "（二）投标人认为需提供的其他资料", spec, bold=True, space_before=4,
                   space_after=2)


# --------------------------------------------------------------------------- #
# Document assembly
# --------------------------------------------------------------------------- #
def build_bid_document(ctx: BidContext) -> Document:
    spec = ctx.spec
    doc = Document()
    _init_styles(doc, spec)
    _apply_layout(doc.sections[0], spec)

    # Cover (its own section, no header / footer).
    _add_paragraph(doc, "", spec, space_after=0)
    _add_paragraph(doc, "投标文件", spec, font=spec.heading_font, size=spec.cover_title_size,
                   bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=120, space_after=60)
    _add_paragraph(doc, f"项目名称：{ctx.fields.get('project_name') or '______'}", spec,
                   size=spec.cover_info_size, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12)
    _add_paragraph(doc, f"项目编号：{ctx.fields.get('project_code') or '______'}", spec,
                   size=spec.cover_info_size, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12)
    _add_paragraph(doc, f"投标人：{config.value(ctx.supplier, 'company_name')}", spec,
                   size=spec.cover_info_size, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12)
    _add_paragraph(doc, "【第____包】", spec, size=spec.cover_info_size,
                   align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12)
    _add_paragraph(doc, config.value(ctx.supplier, "date"), spec, size=spec.cover_info_size,
                   align=WD_ALIGN_PARAGRAPH.CENTER)

    # Content section onwards.
    content_section = doc.add_section(WD_SECTION.NEW_PAGE)
    _apply_layout(content_section, spec)
    _set_header_footer(content_section, spec)

    numbered = _setup_heading_numbering(doc, spec)
    _add_toc(doc, spec)

    toc_items = ctx.template.toc_items
    fill_values = fill.build_values(ctx.fields, ctx.supplier)
    for position, item in enumerate(toc_items, start=1):
        title = _TOC_PREFIX_RE.sub("", item.strip())
        if not numbered:
            title = item.strip()  # keep the manual "一、" prefix as a fallback
        _add_heading(doc, title, 1, spec, numbered=numbered)

        fmt = ctx.template.format_for(position)
        if fmt is not None and fmt.title:
            # 每个格式的题名：居中，上下各一个空行（与招标文件模板一致）。
            _add_paragraph(doc, "", spec, space_before=0, space_after=0)
            _add_paragraph(doc, fmt.title, spec, font=spec.heading_font, size=spec.h2_size,
                           bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
                           space_before=0, space_after=0)
            _add_paragraph(doc, "", spec, space_before=0, space_after=0)
        if position == _MATERIALS_POSITION:
            _materials_section(doc, ctx, spec)
            continue
        if fmt is not None:
            lines = fmt.body
            pending = _placed_tables(fmt, position)
            for i, line in enumerate(lines):
                if not line.text.strip():
                    continue
                # flush template tables that sit above this line in the tender
                while pending and pending[0][0] < (line.page, line.top):
                    _add_table(doc, pending.pop(0)[1], spec)
                    _add_paragraph(doc, "", spec, space_after=4)
                break_before = False
                if i > 0:
                    prev = lines[i - 1]
                    if line.own_page and not prev.own_page:
                        break_before = True        # a self-contained block starts here
                    elif line.own_page and prev.own_page and prev.page != line.page:
                        break_before = True        # a different self-contained page follows
                    elif not line.own_page and prev.own_page:
                        break_before = True        # the self-contained block ends here
                align = _ALIGN_BY_NAME.get(line.align, WD_ALIGN_PARAGRAPH.JUSTIFY)
                indent = spec.first_line_indent_chars if align == WD_ALIGN_PARAGRAPH.JUSTIFY else 0
                _add_paragraph(doc, fill.fill_marked(line.text, fill_values), spec,
                               align=align, indent_chars=indent, bold=line.bold,
                               line_spacing=spec.body_line_spacing, space_after=2,
                               page_break_before=break_before)
            for _, table_rows in pending:
                _add_table(doc, table_rows, spec)
                _add_paragraph(doc, "", spec, space_after=4)
        if position == _DEVIATION_POSITION:
            _deviation_table(doc, ctx, spec, fmt)

    return doc
